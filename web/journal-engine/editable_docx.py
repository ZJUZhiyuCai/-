"""Export the current Shuye model as editable OOXML, not page screenshots.

Source OMML remains native Office Math. Unsupported models fail explicitly.
"""
import base64
from copy import deepcopy
import hashlib
from io import BytesIO
import re
from lxml import etree as ET
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from journal import prepare_blocks, ROLE_STYLE, FRONT_ROLES

W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
NS={'w':W,'m':M}

def element(tag, **attrs):
    e=OxmlElement(tag)
    for k,v in attrs.items():e.set(qn('w:'+k),str(v))
    return e

def field(p, instruction, value='1'):
    p._p.append(element('w:r'));r=p._p[-1];r.append(element('w:fldChar',fldCharType='begin'))
    r=element('w:r');e=element('w:instrText');e.text=' '+instruction+' ';e.set('{http://www.w3.org/XML/1998/namespace}space','preserve');r.append(e);p._p.append(r)
    r=element('w:r');r.append(element('w:fldChar',fldCharType='separate'));p._p.append(r)
    p.add_run(str(value));r=element('w:r');r.append(element('w:fldChar',fldCharType='end'));p._p.append(r)

def bookmark(p,name,ident):
    p._p.append(element('w:bookmarkStart',id=ident,name=name));p._p.append(element('w:bookmarkEnd',id=ident))

def font(run,size=10.5,bold=None,family='serif'):
    run.font.name='FreeSerif';run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC' if family=='sans' else 'Noto Serif CJK SC')

class Writer:
    def __init__(self,plan):
        self.doc=Document();self.plan=plan;self.counts={'textBlocks':0,'tables':0,'images':0,'math':0,'notes':0}
        self.footnotes=element('w:footnotes');self.note_number=0;self.note_map={}
        for n,typ in [(-1,'separator'),(0,'continuationSeparator')]:
            e=element('w:footnote',id=n,type=typ);p=element('w:p');r=element('w:r');r.append(element('w:'+typ));p.append(r);e.append(p);self.footnotes.append(e)
        normal=self.doc.styles['Normal'];normal.font.name='FreeSerif';normal.font.size=Pt(10.5)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Serif CJK SC')
        normal.paragraph_format.space_after=Pt(0)
        for name in ['Title','Heading 1','Heading 2','Heading 3','Heading 4']:
            style=self.doc.styles[name];style.font.color.rgb=RGBColor(0,0,0)
            color=style.element.get_or_add_rPr().find(qn('w:color'))
            if color is not None:
                for key in ('themeColor','themeTint','themeShade'):color.attrib.pop(qn('w:'+key),None)
            for borders in style.element.findall('.//'+qn('w:pBdr')):borders.getparent().remove(borders)
        self.doc.settings.element.append(element('w:updateFields',val='true'))
        self.doc.settings.odd_and_even_pages_header_footer=True
        self.setup(self.doc.sections[0],{},None)

    def setup(self,section,layout,page_start):
        p=self.plan;section.page_width=Mm(p['width']);section.page_height=Mm(p['height'])
        section.top_margin=Mm(layout.get('top',22));section.bottom_margin=Mm(layout.get('bottom',22));section.left_margin=Mm(layout.get('inner',22));section.right_margin=Mm(layout.get('outer',20))
        section.header_distance=Mm(10);section.footer_distance=Mm(10)
        cols=section._sectPr.find(qn('w:cols'))
        if cols is None:cols=element('w:cols');section._sectPr.append(cols)
        cols.set(qn('w:num'),str(layout.get('columns',1)));cols.set(qn('w:space'),str(round(layout.get('gap',8)*1440/25.4)))
        for header,key in [(section.header,'header'),(section.even_page_header,'headerEven')]:
            header.is_linked_to_previous=False;h=header.paragraphs[0];h.text=layout.get(key) or layout.get('header') or p['title'];h.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in h.runs:font(r,8)
        for footer,key in [(section.footer,'footer'),(section.even_page_footer,'footerEven')]:
            footer.is_linked_to_previous=False;paragraph=footer.paragraphs[0];paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
            template=layout.get(key) or layout.get('footer') or '{page}'
            parts=template.split('{page}')
            for i,part in enumerate(parts):
                if i:field(paragraph,'PAGE',str(page_start or 1))
                font(paragraph.add_run(part),8)
        pg=section._sectPr.find(qn('w:pgNumType'))
        if pg is not None:section._sectPr.remove(pg)
        if page_start is not None:section._sectPr.append(element('w:pgNumType',start=page_start))

    def formula(self,p,model):
        source=model.get('sourceOmml')
        if not source or hashlib.sha256(source.encode()).hexdigest()!=model.get('sourceSha256'):raise ValueError('EDITABLE_OMML_SOURCE_REQUIRED: '+model.get('id',''))
        root=ET.fromstring(source.encode(),ET.XMLParser(resolve_entities=False,no_network=True))
        if root.tag not in ('{'+M+'}oMath','{'+M+'}oMathPara'):raise ValueError('INVALID_OMML_ROOT')
        if root.tag=='{'+M+'}oMathPara':
            for child in root:
                if child.tag=='{'+M+'}oMath':p._p.append(deepcopy(child))
        else:p._p.append(root)
        self.counts['math']+=1

    def note(self,p,ident):
        if ident not in self.notes:raise ValueError('NOTE_NOT_FOUND')
        if ident not in self.note_map:
            self.note_number+=1;self.note_map[ident]=self.note_number;n=self.notes[ident]
            node=element('w:footnote',id=self.note_number);paragraph=element('w:p');pr=element('w:pPr');pr.append(element('w:pStyle',val='FootnoteText'));paragraph.append(pr)
            r=element('w:r')
            if n.get('marker'):t=element('w:t');t.text=n['marker'];r.append(t)
            else:r.append(element('w:footnoteRef'))
            paragraph.append(r)
            r=element('w:r');t=element('w:t');t.text=' '+n['text'];r.append(t);paragraph.append(r);node.append(paragraph);self.footnotes.append(node);self.counts['notes']+=1
        r=element('w:r');pr=element('w:rPr');pr.append(element('w:vertAlign',val='superscript'));r.append(pr);ref=element('w:footnoteReference',id=self.note_map[ident]);r.append(ref)
        if self.notes[ident].get('marker'):
            ref.set(qn('w:customMarkFollows'),'1');t=element('w:t');t.text=self.notes[ident]['marker'];r.append(t)
        p._p.append(r)

    def inline(self,p,text,block,marks=None):
        formulas={m['id']:m for m in block.get('inlineMath',[])};marks=block.get('marks',[]) if marks is None else marks
        # Python offsets are Unicode codepoints; persisted mark offsets are UTF-16.
        def offset(i):return len(text[:i].encode('utf-16-le'))//2
        tokens=list(re.finditer(r'\[\[math:([-\w]+)\]\]|\[\^([-\w]+)\]|\[\[(ref|cite):([-\w]+)\]\]',text))
        def plain(start,end):
            if start>=end:return
            points=[start]+[i for i in range(start+1,end) if any(offset(i) in (m['start'],m['end']) for m in marks)]+[end]
            for a,b in zip(points,points[1:]):
                style={}
                for m in marks:
                    if m['start']<=offset(a)<m['end']:style.update(m)
                run=p.add_run(text[a:b]);font(run,style.get('fontSize',self.size),style.get('bold',self.bold),style.get('fontFamily','serif'))
                for key in ('italic','underline'):setattr(run,key,style.get(key,False))
                run.font.superscript=style.get('superscript',False);run.font.subscript=style.get('subscript',False)
                if style.get('color'):run.font.color.rgb=RGBColor.from_string(style['color'][1:])
        start=0
        for m in tokens:
            plain(start,m.start())
            if m[1]:self.formula(p,formulas[m[1]])
            elif m[2]:self.note(p,m[2])
            elif m[3]=='cite':
                ids=list(self.bibliography)
                if m[4] not in ids:raise ValueError('BIBLIOGRAPHY_NOT_FOUND')
                p.add_run('['+str(ids.index(m[4])+1)+']')
            else:
                if m[4] not in self.by_id:raise ValueError('REFERENCE_NOT_FOUND')
                p.add_run(self.by_id[m[4]].get('text',''))
            start=m.end()
        plain(start,len(text))

    def picture(self,p,b,width_mm):
        match=re.fullmatch(r'data:image/(png|jpeg|jpg|webp);base64,(.+)',b['src'],re.S)
        if not match:raise ValueError('UNSUPPORTED_WORD_IMAGE')
        data=base64.b64decode(match[2],validate=True)
        if match[1]=='webp':
            from PIL import Image
            out=BytesIO();Image.open(BytesIO(data)).save(out,'PNG');data=out.getvalue()
        from PIL import Image
        with Image.open(BytesIO(data)) as im:w,h=im.size
        height_mm=self.plan['height']-self.layout.get('top',22)-self.layout.get('bottom',22)-(35 if b.get('captionIds') else 15)
        width_mm=min(width_mm,height_mm*w/h)
        p.add_run().add_picture(BytesIO(data),width=Mm(width_mm));p.alignment=WD_ALIGN_PARAGRAPH.CENTER;self.counts['images']+=1

    def block(self,b,parent=None):
        parent=parent or self.doc;typ=b['type'];role=b.get('role');self.size=b.get('fontSize',self.layout.get('fontSize',10.5));self.bold=b.get('bold',False)
        leading=self.size*self.layout.get('lineHeight',1.5);align=b.get('align','justify')
        if role in ROLE_STYLE:
            self.size,leading,align,self.bold=ROLE_STYLE[role];custom=self.journal.get('roleStyles',{}).get(role,{})
            self.size=custom.get('fontSize',self.size);leading=custom.get('leading',leading);self.bold=custom.get('bold',self.bold)
        if typ=='pageBreak':parent.add_paragraph().add_run().add_break(__import__('docx.enum.text',fromlist=['WD_BREAK']).WD_BREAK.PAGE);return
        if typ=='table':
            rows=b['rows'];table=parent.add_table(rows=len(rows),cols=len(rows[0]));table.autofit=False
            widths=b.get('columnWidths',[1]*len(rows[0]));total=sum(widths);usable=self.usable
            for col,w in zip(table.columns,widths):col.width=Mm(usable*w/total)
            table.style='Table Grid' if b.get('borderStyle','grid')=='grid' else 'Normal Table'
            if b.get('borderStyle')=='three-line':
                borders=element('w:tblBorders')
                for edge in ('top','bottom','left','right','insideH','insideV'):borders.append(element('w:'+edge,val='single' if edge in ('top','bottom') else 'nil',sz=8))
                table._tbl.tblPr.append(borders)
            for ri,row in enumerate(rows):
                if ri<b.get('headerRows',0):table.rows[ri]._tr.get_or_add_trPr().append(element('w:tblHeader'))
                for ci,value in enumerate(row):
                    cell=table.cell(ri,ci);cell.width=Mm(usable*widths[ci]/total);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.1
                    if b.get('borderStyle')=='three-line' and ri==b.get('headerRows',1)-1:
                        borders=element('w:tcBorders');borders.append(element('w:bottom',val='single',sz=4));cell._tc.get_or_add_tcPr().append(borders)
                    marks=next((x['marks'] for x in b.get('cellMarks',[]) if (x['row'],x['col'])==(ri,ci)),[])
                    pictures=[entry for entry in b.get('cellImages',[]) if (entry['row'],entry['col'])==(ri,ci)]
                    for k,entry in enumerate(pictures):
                        if k:p=cell.add_paragraph()
                        self.picture(p,self.by_id[entry['imageId']],max(5,usable*widths[ci]/total-4));p.paragraph_format.keep_with_next=bool(value) or k<len(pictures)-1
                    if pictures:p=cell.add_paragraph();p.paragraph_format.line_spacing=1.1
                    self.inline(p,value,b,marks)
            for merge in b.get('merges',[]):
                r,c=merge['row'],merge['col'];table.cell(r,c).merge(table.cell(r+merge['rowspan']-1,c+merge['colspan']-1))
            self.counts['tables']+=1
            if b.get('text'):self.block({'type':'paragraph','text':b['text'],'fontSize':9})
        else:
            p=parent.add_paragraph();p.paragraph_format.line_spacing=Pt(leading);p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST;p.paragraph_format.widow_control=True
            if b.get('_keepWithNext') or b.get('captionIds'):p.paragraph_format.keep_with_next=True
            p.alignment={'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT,'left':WD_ALIGN_PARAGRAPH.LEFT,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}.get(align,WD_ALIGN_PARAGRAPH.JUSTIFY)
            if typ=='heading':p.style='Heading '+str(min(4,b.get('level',1)));p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(7)
            if typ in ('paragraph','heading'):
                self.inline(p,b.get('text',''),b);self.counts['textBlocks']+=1
                if not role and typ=='paragraph':p.paragraph_format.first_line_indent=Pt(b.get('indentPt',b.get('indent',0)*self.size))
            elif typ=='equation':
                self.formula(p,b);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                if b.get('equationNumber'):p.add_run('  ('+b['equationNumber']+')')
                if b.get('equationSuffix'):p.add_run(b['equationSuffix'])
            elif typ=='image':self.picture(p,b,self.usable*b.get('width',100)/100)
            else:raise ValueError('WORD_BLOCK_UNSUPPORTED: '+typ)
        captions=b.get('captionIds',[])
        if captions and typ=='table':
            for cell in table.rows[-1].cells:
                for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
        for i,ident in enumerate(captions):self.block({**self.by_id[ident],'_keepWithNext':i<len(captions)-1},parent)

    def article(self,record,index,first_page=None):
        source=record['document'];self.layout=source.get('layout',{});self.journal=source.get('journal',{});self.notes={n['id']:n for n in source.get('notes',[])};self.note_map={};self.bibliography={n['id']:n for n in source.get('bibliography',[])}
        if source.get('composition'):raise ValueError('WORD_COMPOSITION_UNSUPPORTED')
        if self.layout.get('writingMode','horizontal-tb')!='horizontal-tb':raise ValueError('WORD_VERTICAL_UNSUPPORTED')
        if self.layout.get('notesMode','footnote')!='footnote':raise ValueError('WORD_ENDNOTE_MODE_UNSUPPORTED')
        if self.plan['width']-self.layout.get('inner',22)-self.layout.get('outer',20)-(self.layout.get('columns',1)-1)*self.layout.get('gap',8)<30:raise ValueError('WORD_TEXT_AREA_TOO_NARROW')
        if any(b.get('ruby') or b.get('sideTextIds') for b in source['blocks']):raise ValueError('WORD_RUBY_SIDE_GROUP_UNSUPPORTED')
        blocks,self.by_id=prepare_blocks(source)
        section=self.doc.sections[0] if index==0 else self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        self.setup(section,{**self.layout,'columns':1},first_page)
        self.usable=self.plan['width']-self.layout.get('inner',22)-self.layout.get('outer',20)
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.paragraph_format.keep_with_next=True;bookmark(p,'article_'+str(index),index+1)
        in_body=False
        for b in blocks:
            if b.get('role') not in FRONT_ROLES and not in_body:
                in_body=True
                if self.layout.get('columns',1)>1:
                    section=self.doc.add_section(WD_SECTION_START.CONTINUOUS);self.setup(section,self.layout,None)
                    self.usable=(self.usable-self.layout.get('gap',8)*(self.layout['columns']-1))/self.layout['columns']
            self.block(b)
        for i,n in enumerate(self.bibliography.values(),1):self.block({'type':'paragraph','text':'['+str(i)+'] '+n['text']})

    def save(self,path):
        if self.counts['notes']:
            part=Part(PackURI('/word/footnotes.xml'),'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',ET.tostring(self.footnotes,xml_declaration=True,encoding='UTF-8'),self.doc.part.package)
            self.doc.part.relate_to(part,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes')
        self.doc.save(path)
        return dict(self.counts)

def export_article(record,plan,path,first_page):
    writer=Writer(plan);writer.article(record,0,first_page);return writer.save(path)
