"""Version-bound issue composition, measured contents and editable delivery.

The InkQ workflow is integrated using Shuye's own documents and storage. No
InkQ accounts, credit ledger, private source packages or customer data are copied.
"""
import hashlib
import json
from pathlib import Path
from io import BytesIO
import zipfile
import re
from collections import Counter
from copy import deepcopy
from lxml import etree as ET
from pypdf import PdfReader
from docx.shared import Pt
from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
from editable_docx import Writer,export_article,field,font,NS
from word_render import render
from proof import text_geometry
from omml_source import validated_source

def sha(data):return hashlib.sha256(data).hexdigest()
def title(article,entry):
    return re.sub(r'\[\^[-\w]+\]','',entry.get('title') or next((b['text'] for b in article['document']['blocks'] if b.get('role')=='title-zh'),article['document']['title']))

def source_counts(document):
    return {'textBlocks':sum(b['type'] in ('paragraph','heading') for b in document['blocks']),
            'tables':sum(b['type']=='table' for b in document['blocks']),
            'images':sum(b['type']=='image' for b in document['blocks']),
            'math':sum((1 if b['type']=='equation' else 0)+len(b.get('inlineMath',[])) for b in document['blocks'])}

def word_contents(path,values):
    """Cache measured page references while retaining editable PAGEREF fields."""
    with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    root=ET.fromstring(parts['word/document.xml']);namespace=NS['w']
    for p in root.findall('.//w:p',NS):
        instructions=p.findall('.//w:instrText',NS)
        if not instructions:continue
        instruction=''.join(e.text or '' for e in instructions).strip()
        if instruction.startswith('PAGEREF article_'):
            key=instruction.split()[1]
            if key not in values:raise ValueError('CONTENTS_BOOKMARK_MISSING')
            after=False
            for r in p.findall('w:r',NS):
                marker=r.find('w:fldChar',NS)
                if marker is not None:
                    if marker.get('{'+namespace+'}fldCharType')=='separate':after=True
                    if marker.get('{'+namespace+'}fldCharType')=='end':break
                if after:
                    texts=r.findall('w:t',NS)
                    if texts:texts[0].text=str(values[key]);break
    parts['word/document.xml']=ET.tostring(root,xml_declaration=True,encoding='UTF-8')
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)

def missing_figures(document):
    import re
    owners={c for b in document['blocks'] if b['type']=='image' or b['type']=='table' and b.get('cellImages') for c in b.get('captionIds',[])}
    errors=[];blocks=document['blocks']
    for i,b in enumerate(blocks):
        label=re.match(r'^图\s*(\d+)(.*)',b.get('text','').strip());following=re.match(r'^Fig\.?\s*(\d+)\b',blocks[i+1].get('text','').strip(),re.I) if i+1<len(blocks) else None
        if b['type'] in ('paragraph','heading') and label and (re.match(r'^(?:\s|[.:：])',label[2]) or following and following[1]==label[1]) and b['id'] not in owners:errors.append({'code':'FIGURE_ASSET_UNASSIGNED','severity':'error','blockId':b['id']})
    return errors

def verify_contents(contents,articles):
    compact=re.sub(r'\s+','',contents);cursor=0
    for article in articles:
        heading=re.sub(r'\s+','',article['title']);start=compact.find(heading,cursor)
        if start<0:raise ValueError('CONTENTS_TITLE_NOT_RENDERED')
        cursor=start+len(heading);number=re.match(r'[.·…]*([0-9]+)',compact[cursor:])
        if not number or int(number[1])!=article['firstPage']:raise ValueError('CONTENTS_PAGE_REFERENCE_NOT_RENDERED')
        cursor+=number.end()

def verify_docx(path,records):
    """Verify preserved source math XML and editable text, including cell text."""
    with zipfile.ZipFile(path) as z:
        root=ET.fromstring(z.read('word/document.xml'))
        def paragraph_text(p):
            return ''.join(e.text or '' if e.tag=='{'+NS['w']+'}t' else '\t' if e.tag=='{'+NS['w']+'}tab' else '\n' for e in p.iter() if e.tag in ('{'+NS['w']+'}t','{'+NS['w']+'}tab','{'+NS['w']+'}br'))
        paragraphs=Counter(paragraph_text(p) for p in root.findall('.//w:p',NS));checked=0
        footnotes=ET.fromstring(z.read('word/footnotes.xml')) if 'word/footnotes.xml' in z.namelist() else ET.Element('empty')
        note_texts=Counter(paragraph_text(p).strip() for p in footnotes.findall('.//w:p',NS))
        for record in records:
            document=record['document'];notes={n['id']:n for n in document.get('notes',[])};blocks={b['id']:b for b in document['blocks']};bib=[n['id'] for n in document.get('bibliography',[])]
            def displayed(text):
                def replace(m):
                    if m[1]:return ''
                    if m[2]:return notes[m[2]].get('marker','')
                    return '['+str(bib.index(m[4])+1)+']' if m[3]=='cite' else blocks[m[4]].get('text','')
                return re.sub(r'\[\[math:([-\w]+)\]\]|\[\^([-\w]+)\]|\[\[(ref|cite):([-\w]+)\]\]',replace,text)
            for b in document['blocks']:
                values=([b.get('text','')] if b['type'] in ('paragraph','heading','table') else [])+([v for row in b['rows'] for v in row] if b['type']=='table' else [])
                for value in values:
                    value=displayed(value)
                    if not value:continue
                    if paragraphs[value]<=0:raise ValueError('WORD_TEXT_SOURCE_MISMATCH: '+b['id'])
                    paragraphs[value]-=1;checked+=1
            for note in notes.values():
                value=((note.get('marker') or '')+' '+note['text']).strip()
                if note_texts[value]<=0:raise ValueError('WORD_NOTE_SOURCE_MISMATCH: '+note['id'])
                note_texts[value]-=1
        # Namespace prefixes differ between serializers; compare expanded-name trees.
        def tree(e):return (e.tag,tuple(sorted(e.attrib.items())),e.text or '',tuple(tree(c) for c in e))
        actual=[tree(m) for m in root.findall('.//m:oMath',NS)];expected=[]
        for r in records:
            for b in r['document']['blocks']:
                for m in ([b] if b['type']=='equation' else b.get('inlineMath',[])):
                    source=validated_source(m);expected.extend(tree(e) for e in ([source] if source.tag.endswith('}oMath') else source.findall('m:oMath',NS)))
        remaining=list(actual)
        for e in expected:
            if e not in remaining:raise ValueError('WORD_FORMULA_SOURCE_MISMATCH')
            remaining.remove(e)
        if remaining:raise ValueError('WORD_FORMULA_DUPLICATED')
        return {'nativeMathElements':len(actual),'formulaSourceExact':True,'editableTextSegments':checked,'textSourceExact':True,'editable':True,'imageParts':sum(n.startswith('word/media/') for n in z.namelist())}

def render_issue(request):
    folder=Path(request['directory']);input=request['document'];plan=input['issuePlan'];records=input['articles'];entries=plan['articles']
    if len(records)!=len(entries):raise ValueError('ISSUE_SOURCE_MISMATCH')
    writer=Writer(plan)
    p=writer.doc.add_paragraph(plan['title'],style='Title')
    for r in p.runs:font(r,18,True)
    writer.doc.add_paragraph('目录')
    for i,(record,entry) in enumerate(zip(records,entries),1):
        p=writer.doc.add_paragraph();p.paragraph_format.space_after=Pt(8);p.paragraph_format.tab_stops.add_tab_stop(Pt((plan['width']-42)*72/25.4),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        p.add_run(title(record,entry)+'\t');field(p,'PAGEREF article_'+str(i)+' \\h',str(plan['firstPage']))
        authors=entry.get('authors') or next((b['text'] for b in record['document']['blocks'] if b.get('role')=='authors-zh'),'')
        if authors:writer.doc.add_paragraph(re.sub(r'\[\^[-\w]+\]','',authors))
    for i,r in enumerate(records,1):writer.article(r,i,plan['firstPage'] if i==1 else None)
    counts=writer.save(folder/'issue.docx');word_audit=verify_docx(folder/'issue.docx',records)
    measured=render(folder/'issue.docx',folder/'issue.pdf');reader=PdfReader(folder/'issue.pdf');total=len(reader.pages)
    starts=[measured['bookmarks']['article_'+str(i)] for i in range(1,len(records)+1)]
    if sorted(set(starts))!=starts:raise ValueError('ISSUE_ARTICLE_ORDER_INVALID')
    logical={k:v-starts[0]+plan['firstPage'] for k,v in measured['bookmarks'].items()}
    word_contents(folder/'issue.docx',logical)
    # The cached fields must not change the final pagination.
    final=render(folder/'issue.docx',folder/'issue.pdf')
    if final['bookmarks']!=measured['bookmarks']:raise ValueError('ISSUE_PAGINATION_NOT_STABLE')
    word_audit=verify_docx(folder/'issue.docx',records)
    reader=PdfReader(folder/'issue.pdf');total=len(reader.pages)
    articles=[];artifacts=[]
    def artifact(ident,filename,mime):
        data=(folder/filename).read_bytes();a={'id':ident,'filename':filename,'mimeType':mime,'bytes':len(data),'sha256':sha(data)};artifacts.append(a);return a
    docx_mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    for i,(record,entry,start) in enumerate(zip(records,entries,starts),1):
        first=start-starts[0]+plan['firstPage'];last=(starts[i]-1 if i<len(starts) else total)-starts[0]+plan['firstPage']
        path=folder/f'article-{i}.docx';c=export_article(record,plan,path,first);audit=verify_docx(path,[record]);expected=source_counts(record['document'])
        for key,value in expected.items():
            if c[key]!=value:raise ValueError('WORD_OBJECT_COUNT_MISMATCH: '+key)
        errors=missing_figures(record['document']);articles.append({'projectId':record['projectId'],'revision':record['revision'],'title':title(record,entry),'firstPage':first,'lastPage':last,'physicalFirstPage':start,'pages':last-first+1,'counts':c,'word':audit,'quality':{'status':'BLOCKED' if errors else 'REVIEW','diagnostics':errors}})
        artifact('article-'+str(i),path.name,docx_mime)
    geometry=text_geometry((folder/'issue.pdf').read_bytes())
    if geometry['status']!='COMPLETE':raise ValueError('ISSUE_GEOMETRY_INCOMPLETE')
    errors=[{**d,'projectId':a['projectId']} for a in articles for d in a['quality']['diagnostics']]+geometry['outside']
    parse_errors=sum(line['text'].count('¿') for page in geometry['pages'] for line in page['lines'])
    source_errors=json.dumps(records,ensure_ascii=False).count('¿')
    if parse_errors>source_errors:errors.append({'code':'WORD_FORMULA_PARSE_ERROR','severity':'error','count':parse_errors})
    # Verify contents references appear in the rendered contents, not only XML.
    contents=''.join(reader.pages[i].extract_text() or '' for i in range(starts[0]-1))
    verify_contents(contents,articles)
    report={'scope':'complete-issue','sourceDocumentSha256':request['sourceDocumentSha256'],'sourceRevision':request['revision'],'pages':total,'contentsPages':starts[0]-1,'articles':articles,'word':word_audit,'counts':counts,'renderer':measured['renderer'],'contentsRecalculated':True,'paginationStable':True,'quality':{'status':'BLOCKED' if errors else 'REVIEW','productionReady':False,'diagnostics':errors,'notChecked':['编辑逐字与公式语义签样','完整 PDF/X 与印厂接收']},'geometry':{'status':geometry['status'],'characters':geometry['characters'],'outside':geometry['outside']}}
    report['wordRendering']=final['compatibility'];report['wordRendering']['parseErrorGlyphs']=parse_errors
    (folder/'report.json').write_text(json.dumps(report,ensure_ascii=False),encoding='utf-8')
    artifact('issue-pdf','issue.pdf','application/pdf');artifact('issue-word','issue.docx',docx_mime);artifact('report','report.json','application/json')
    with zipfile.ZipFile(folder/'delivery.zip','w',zipfile.ZIP_DEFLATED) as z:
        for a in artifacts:z.write(folder/a['filename'],a['filename'])
    artifact('delivery','delivery.zip','application/zip')
    return {'pages':total,'artifacts':artifacts,'pdfSha256':sha((folder/'issue.pdf').read_bytes()),'reportSha256':sha((folder/'report.json').read_bytes())}
